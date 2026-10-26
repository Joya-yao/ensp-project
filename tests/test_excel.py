"""地址表 Excel 的写入、读回与覆盖规则。"""

from pathlib import Path

import pytest

from labgen.excel import (
    DEVICE_SHEET,
    TERMINAL_SHEET,
    apply_overrides,
    available,
    read_address_xlsx,
    write_address_xlsx,
)
from labgen.requirements import load_requirements
from labgen.topo import parse_topo

pytestmark = pytest.mark.skipif(not available(), reason="没装 openpyxl")

ROOT = Path(__file__).resolve().parents[1]
LAB = ROOT / "docs" / "labs"


@pytest.fixture(scope="module")
def topo():
    return parse_topo(LAB / "lab01.topo")


@pytest.fixture(scope="module")
def req():
    return load_requirements(LAB / "lab01-vlan.yaml")


def test_workbook_has_two_sheets(tmp_path, topo, req):
    from openpyxl import load_workbook

    path = write_address_xlsx(tmp_path / "a.xlsx", topo, req)
    book = load_workbook(path)
    assert DEVICE_SHEET in book.sheetnames
    assert TERMINAL_SHEET in book.sheetnames


def test_roundtrip_keeps_addresses(tmp_path, topo, req):
    path = write_address_xlsx(tmp_path / "a.xlsx", topo, req)
    sheet = read_address_xlsx(path)
    assert sheet["devices"][0]["设备"] == "LSW1"
    assert sheet["devices"][0]["IP"] == "10.10.0.254"
    assert len(sheet["terminals"]) == 4
    assert sheet["terminals"][0]["接在"].startswith("LSW1:")


def test_skeleton_when_no_requirements(tmp_path, topo):
    path = write_address_xlsx(tmp_path / "b.xlsx", topo, None)
    sheet = read_address_xlsx(path)
    assert sheet["devices"], "骨架表也要列出设备接口"
    assert sheet["devices"][0]["IP"] == ""


def test_device_ip_override(tmp_path, topo, req):
    path = write_address_xlsx(tmp_path / "c.xlsx", topo, req)
    sheet = read_address_xlsx(path)
    sheet["devices"][0]["IP"] = "10.10.0.99"
    sheet["devices"][0]["掩码"] = "255.255.255.0"

    patched = apply_overrides(req, sheet)
    iface = next(i for i in patched.devices["LSW1"].interfaces if i.name == "Vlanif10")
    assert iface.ip == "10.10.0.99/24"


def test_blank_value_keeps_yaml(tmp_path, topo, req):
    path = write_address_xlsx(tmp_path / "d.xlsx", topo, req)
    sheet = read_address_xlsx(path)
    sheet["devices"][0]["IP"] = "?"
    patched = apply_overrides(req, sheet)
    iface = next(i for i in patched.devices["LSW1"].interfaces if i.name == "Vlanif10")
    assert iface.ip == "10.10.0.254/24"


def test_terminal_vlan_override(tmp_path, topo, req):
    path = write_address_xlsx(tmp_path / "e.xlsx", topo, req)
    sheet = read_address_xlsx(path)
    row = next(r for r in sheet["terminals"] if r["终端"] == "PC1")
    row["VLAN"] = "99"

    patched = apply_overrides(req, sheet)
    iface = next(i for i in patched.devices["LSW1"].interfaces if i.name == "GigabitEthernet0/0/1")
    assert iface.vlan == 99


def test_unknown_device_is_ignored(tmp_path, topo, req):
    path = write_address_xlsx(tmp_path / "f.xlsx", topo, req)
    sheet = read_address_xlsx(path)
    sheet["devices"].append({"设备": "AR9", "接口": "GigabitEthernet0/0/0", "IP": "1.1.1.1"})
    patched = apply_overrides(req, sheet)
    assert "AR9" not in patched.devices