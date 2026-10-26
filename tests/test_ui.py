"""本地 UI 的接口测试：起真服务，用 HTTP 打一遍。"""

import base64
import json
import threading
import urllib.error
import urllib.request
from pathlib import Path
from urllib.parse import quote

import pytest

from labgen.ui import make_server

ROOT = Path(__file__).resolve().parents[1]
LAB = ROOT / "docs" / "labs"


@pytest.fixture()
def server(tmp_path):
    srv, url = make_server(
        0,
        root=tmp_path,
        work=tmp_path / "output" / "ui" / "work",
        final=tmp_path / "output" / "ui" / "final",
    )
    thread = threading.Thread(target=srv.serve_forever, daemon=True)
    thread.start()
    try:
        yield srv, url
    finally:
        srv.shutdown()
        thread.join(timeout=2)


def _token(url: str) -> str:
    return url.split("t=")[1]


def _post(url: str, path: str, payload: dict, token: str | None = None):
    body = json.dumps(payload).encode("utf-8")
    request = urllib.request.Request(
        url.split("/?")[0] + path, data=body, method="POST"
    )
    request.add_header("Content-Type", "application/json")
    request.add_header("X-Token", token if token is not None else _token(url))
    try:
        with urllib.request.urlopen(request, timeout=5) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _get(url: str, path: str):
    try:
        with urllib.request.urlopen(url.split("/?")[0] + path, timeout=5) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _yaml():
    return (LAB / "lab01-vlan.yaml").read_text(encoding="utf-8")


def test_index_serves_page_with_token(server):
    srv, url = server
    with urllib.request.urlopen(url, timeout=5) as response:
        html = response.read().decode("utf-8")
    assert response.status == 200
    assert "eNSP 配置生成器" in html
    assert "__TOKEN__" not in html
    assert _token(url) in html


def test_api_rejects_missing_token(server):
    srv, url = server
    body = json.dumps({"topo": "x"}).encode()
    request = urllib.request.Request(
        url.split("/?")[0] + "/api/summary", data=body, method="POST"
    )
    request.add_header("Content-Type", "application/json")
    with pytest.raises(urllib.error.HTTPError) as exc:
        urllib.request.urlopen(request, timeout=5)
    assert exc.value.code == 403


def test_summary_needs_topo(server):
    srv, url = server
    status, data = _post(url, "/api/summary", {})
    assert status == 400
    assert data["ok"] is False
    assert ".topo" in data["error"]


def test_validate_reports_error(server):
    srv, url = server
    status, data = _post(url, "/api/validate", {"yaml": "devices: {}"})
    assert status == 400
    assert data["ok"] is False


def test_non_string_field_is_rejected(server):
    srv, url = server
    status, data = _post(url, "/api/generate", {"yaml": {"a": 1}, "topo": "x.topo"})
    assert status == 400
    assert "字符串" in data["error"]


def test_address_table_saves_markdown_to_work(server, tmp_path):
    srv, url = server
    status, data = _post(
        url,
        "/api/address-table",
        {"yaml": _yaml(), "topo": str(LAB / "lab01.topo"), "work_dir": "output/ui/work"},
    )
    assert data["ok"] is True
    assert "## 网段与网关" in data["markdown"]
    # 文件名带拓扑前缀，便于多项目区分
    assert data["stem"] == "lab01"
    suffix = ".xlsx" if data["xlsx"] else ".md"
    assert data["name"] == f"lab01-addresses{suffix}"
    saved = tmp_path / "output" / "ui" / "work" / f"lab01-addresses{suffix}"
    assert saved.is_file()
    assert data["saved"] == str(saved.resolve())
    assert data["mode"] == "yaml"


def test_address_table_skeleton_without_yaml(server, tmp_path):
    srv, url = server
    status, data = _post(
        url,
        "/api/address-table",
        {"yaml": "", "topo": str(LAB / "lab01.topo"), "work_dir": "output/ui/work"},
    )
    assert data["ok"] is True, data
    assert data["mode"] == "skeleton"
    assert "待填写" in data["markdown"]
    assert "| LSW1 |" in data["markdown"]
    suffix = ".xlsx" if data["xlsx"] else ".md"
    assert (tmp_path / "output" / "ui" / "work" / f"lab01-addresses{suffix}").is_file()


def test_address_table_requires_topo_or_yaml(server):
    srv, url = server
    status, data = _post(url, "/api/address-table", {"yaml": "", "topo": ""})
    assert status == 400
    assert ".topo" in data["error"]


def test_address_table_honours_custom_path(server, tmp_path):
    srv, url = server
    status, data = _post(
        url,
        "/api/address-table",
        {"yaml": _yaml(), "addr_path": "output/ui/work/地址表-实验一.md"},
    )
    assert data["ok"] is True
    assert data["name"] == "地址表-实验一.md"
    assert (tmp_path / "output" / "ui" / "work" / "地址表-实验一.md").is_file()


def test_address_table_rejects_path_outside_output(server):
    srv, url = server
    status, data = _post(
        url, "/api/address-table", {"yaml": _yaml(), "addr_path": "C:/Windows/Temp/a.md"}
    )
    assert status == 400
    assert "output" in data["error"]


def test_upload_writes_bytes_into_work_dir(server, tmp_path):
    srv, url = server
    raw = (LAB / "lab01.topo").read_bytes()
    status, data = _post(
        url,
        "/api/upload",
        {
            "name": "我的实验.topo",
            "data": base64.b64encode(raw).decode("ascii"),
            "work_dir": "output/ui/work",
        },
    )
    assert data["ok"] is True, data
    target = Path(data["path"])
    assert target.is_file()
    assert target.read_bytes() == raw
    assert target.parent == (tmp_path / "output" / "ui" / "work").resolve()


def test_upload_rejects_bad_base64(server):
    srv, url = server
    status, data = _post(url, "/api/upload", {"name": "a.topo", "data": "not base64!"})
    assert status == 400
    assert "base64" in data["error"]


def test_save_and_read_yaml_roundtrip(server, tmp_path):
    srv, url = server
    status, data = _post(
        url,
        "/api/save",
        {"path": "output/ui/work/requirements.yaml", "content": "experiment: {name: 测试}\n"},
    )
    assert data["ok"] is True
    status, data = _get(url, "/api/read?path=" + quote(data["path"], safe=""))
    assert data["ok"] is True
    assert "测试" in data["content"]


def test_save_outside_output_is_rejected(server):
    srv, url = server
    status, data = _post(
        url, "/api/save", {"path": "config.yaml", "content": "x: 1\n"}
    )
    assert status == 400
    assert "output/" in data["error"]


def test_list_shows_work_files(server, tmp_path):
    srv, url = server
    _post(url, "/api/save", {"path": "output/ui/work/a.yaml", "content": "x: 1\n"})
    status, data = _get(url, "/api/list?dir=" + quote("output/ui/work", safe=""))
    assert data["ok"] is True
    assert [f["name"] for f in data["files"]] == ["a.yaml"]


def test_generate_writes_into_final_dir(server, tmp_path):
    srv, url = server
    status, data = _post(
        url,
        "/api/generate",
        {
            "yaml": _yaml(),
            "topo": str(LAB / "lab01.topo"),
            "final_dir": "output/ui/final",
        },
    )
    assert data["ok"] is True, data
    names = {item["name"] for item in data["files"]}
    assert "lab01-LSW1.txt" in names
    assert "lab01-配置文档.docx" in names
    assert (tmp_path / "output" / "ui" / "final" / "lab01-LSW1.txt").is_file()
    # 每个文件都要带绝对路径，页面才能显示「生成到哪了」
    for item in data["files"]:
        assert Path(item["path"]).is_absolute()
        assert Path(item["path"]).is_file()


def test_save_reports_name_and_size(server):
    srv, url = server
    status, data = _post(
        url, "/api/save", {"path": "output/ui/work/x.yaml", "content": "a: 1\n"}
    )
    assert data["ok"] is True
    assert data["name"] == "x.yaml"
    assert data["size"] == 5


def test_generate_rejects_final_dir_outside_output(server):
    srv, url = server
    status, data = _post(
        url,
        "/api/generate",
        {"yaml": _yaml(), "topo": str(LAB / "lab01.topo"), "final_dir": "C:/Windows/Temp/x"},
    )
    assert status == 400
    assert "output" in data["error"]


def test_ai_cloud_without_key_is_rejected(server):
    srv, url = server
    status, data = _post(
        url,
        "/api/ai-convert",
        {
            "text": "要求：VLAN 10",
            "provider": "cloud",
            "base_url": "https://api.deepseek.com/v1",
            "model": "deepseek-chat",
            "api_key": "",
            "mode": "yaml",
        },
    )
    assert status == 400
    assert "api_key" in data["error"]


def test_ai_rejects_unknown_mode(server):
    srv, url = server
    status, data = _post(
        url,
        "/api/ai-convert",
        {
            "text": "要求",
            "provider": "local",
            "base_url": "http://127.0.0.1:11434/v1",
            "model": "qwen2.5:7b",
            "mode": "excel",
        },
    )
    assert status == 400
    assert "mode" in data["error"]


def test_generate_applies_address_xlsx_overrides(server, tmp_path):
    from labgen.excel import available

    if not available():
        pytest.skip("没装 openpyxl")

    from openpyxl import load_workbook

    srv, url = server
    status, data = _post(
        url,
        "/api/address-table",
        {"yaml": _yaml(), "topo": str(LAB / "lab01.topo"), "addr_path": "output/ui/work/lab01-addresses.xlsx"},
    )
    assert data["xlsx"] is True, data

    book = load_workbook(data["saved"])
    sheet = book["设备地址"]
    for row in sheet.iter_rows(min_row=2):
        if row[1].value == "Vlanif10":
            row[2].value = "10.10.0.99"
            row[3].value = "255.255.255.0"
    book.save(data["saved"])

    status, gen = _post(
        url,
        "/api/generate",
        {
            "yaml": _yaml(),
            "topo": str(LAB / "lab01.topo"),
            "final_dir": "output/ui/final",
            "addr_path": data["saved"],
        },
    )
    assert gen["ok"] is True, gen
    assert gen["overrides"] == data["saved"]

    text = "\n".join(
        p.read_bytes().decode("utf-8") for p in (tmp_path / "output" / "ui" / "final").glob("*.txt")
    )
    assert "ip address 10.10.0.99 255.255.255.0" in text


def test_download_rejects_path_outside_output(server):
    srv, url = server
    with pytest.raises(urllib.error.HTTPError) as exc:
        urllib.request.urlopen(
            url.split("/?")[0] + "/api/file?path=" + quote(str(LAB / "lab01.topo"), safe=""),
            timeout=5,
        )
    assert exc.value.code == 403


def test_download_sets_utf8_filename_for_chinese_names(server, tmp_path):
    """Content-Disposition 里不能塞原始中文，要给 filename* 的 UTF-8 形式。"""

    srv, url = server
    target = tmp_path / "output" / "ui" / "final" / "lab01-配置文档.txt"
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text("hello", encoding="utf-8")

    with urllib.request.urlopen(
        url.split("/?")[0] + "/api/file?path=" + quote(str(target), safe=""), timeout=5
    ) as response:
        assert response.status == 200
        disposition = response.headers.get("Content-Disposition")
        assert "filename*=UTF-8''" in disposition
        assert "%E9%85%8D%E7%BD%AE" in disposition  # 「配置」的 UTF-8 编码
        assert response.read() == b"hello"


def test_open_rejects_path_outside_output(server):
    srv, url = server
    status, data = _get(url, "/api/open?path=" + quote(str(LAB / "lab01.topo"), safe=""))
    assert status == 403
    assert data["ok"] is False


def test_open_missing_file_returns_404(server, tmp_path):
    srv, url = server
    missing = tmp_path / "output" / "ui" / "final" / "nope.txt"
    status, data = _get(url, "/api/open?path=" + quote(str(missing), safe=""))
    assert status == 404