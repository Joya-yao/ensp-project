"""把生成结果写成 Word：按设备分节，命令为等宽块，可整段复制。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

from labgen.addresses import (
    address_entries,
    device_rows,
    network_rows,
    parse_ip,
    terminal_rows,
)
from labgen.pastestyle import to_safe_paste
from labgen.requirements import Requirements
from labgen.topo import Topology

__all__ = ["build_document", "write_docx"]

CODE_FONT = "Consolas"
CODE_SIZE = Pt(9)


def _set_code_font(run) -> None:
    run.font.name = CODE_FONT
    run.font.size = CODE_SIZE
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CODE_FONT)


def _shade(paragraph, fill: str = "F2F2F2") -> None:
    pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def _command_block(doc: Document, commands: list[str]) -> None:
    """每行一个零间距段落：复制到命令行时换行不会丢。"""

    for line in commands:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.paragraph_format.left_indent = Pt(6)
        _shade(paragraph)
        run = paragraph.add_run(line)
        _set_code_font(run)


LAYER_NAMES = {"core": "核心层", "aggregation": "汇聚层", "access": "接入层"}


def _device_table(doc: Document, topo: Topology, req: Requirements) -> None:
    table = doc.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    headers = ("设备", "型号", "类型", "部门", "层级", "可连接口数")
    for cell, title in zip(table.rows[0].cells, headers):
        cell.text = title
    for device in topo.devices:
        spec = req.devices.get(device.name)
        row = table.add_row().cells
        row[0].text = device.name
        row[1].text = device.model
        row[2].text = {"router": "路由器", "switch": "交换机", "pc": "终端"}.get(
            device.device_type, device.device_type
        )
        row[3].text = (spec.department if spec and spec.department else "—")
        row[4].text = (
            LAYER_NAMES.get(spec.layer, spec.layer)
            if spec and spec.layer
            else "—"
        )
        row[5].text = str(len(device.ports))


def _link_table(doc: Document, topo: Topology) -> None:
    if not topo.links:
        doc.add_paragraph("没有连线。")
        return
    table = doc.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    for cell, title in zip(table.rows[0].cells, ("设备 A", "接口 A", "设备 B", "接口 B", "线缆")):
        cell.text = title
    for link in topo.links:
        row = table.add_row().cells
        row[0].text = link.src.device
        row[1].text = link.src.port or "未知"
        row[2].text = link.dst.device
        row[3].text = link.dst.port or "未知"
        row[4].text = link.cable


def _add_table(doc: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, title in zip(table.rows[0].cells, headers):
        cell.text = title
    for values in rows:
        row = table.add_row().cells
        for cell, value in zip(row, values):
            cell.text = value


def _note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(text)
    paragraph.runs[0].font.size = Pt(9)


def _address_section(doc: Document, topo: Topology, req: Requirements) -> None:
    if not address_entries(req):
        return

    doc.add_heading("地址与网关", level=1)
    _note(
        doc,
        "配置终端和排查三层互通时看这一节。网关取值优先级：VRRP 虚拟 IP → DHCP 地址池 → "
        "网段内 .254 地址 → 网段内唯一地址；标「本机」表示该设备自己就是网关，标「—」表示"
        "该网段没有可推断的网关（多为设备互联链路），需要人工确认。"
    )

    networks = network_rows(req)
    if networks:
        doc.add_heading("网段与网关", level=2)
        _add_table(
            doc,
            ("网段", "掩码", "网关", "网关在哪", "该网段设备地址", "终端可用地址"),
            networks,
        )

    devices = device_rows(req)
    if devices:
        doc.add_heading("每台设备的地址", level=2)
        _add_table(
            doc, ("设备", "接口", "IP", "掩码", "VLAN", "网关", "说明"), devices
        )

    terminals = terminal_rows(topo, req)
    if terminals:
        doc.add_heading("终端建议地址", level=2)
        _note(
            doc,
            "按拓扑连线自动推出 PC 挂在哪台交换机的哪个口、属于哪个 VLAN，"
            "并从该 VLAN 网段里按顺序取第一个空地址。粘之前可改，别跟已有设备撞。",
        )
        _add_table(doc, ("终端", "接在", "VLAN", "建议 IP", "掩码", "网关"), terminals)

    layer2 = [
        name
        for name, spec in req.devices.items()
        if not any(parse_ip(iface.ip) for iface in spec.interfaces)
    ]
    if layer2:
        _note(doc, "没有三层地址的设备（二层透传，按接入 VLAN 配终端网关即可）：" + "、".join(layer2))


def build_document(
    topo: Topology, req: Requirements, configs: dict[str, list[str]]
) -> Document:
    """构建文档对象；不写盘，便于测试。"""

    doc = Document()

    title = req.experiment or Path(topo.source).stem
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    summary = doc.add_paragraph(
        f"拓扑文件：{Path(topo.source).name}　设备 {len(topo.devices)} 台　"
        f"连线 {len(topo.links)} 条"
    )
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in summary.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    note = doc.add_paragraph()
    note_run = note.add_run(
        "使用方式：按设备标题逐节复制命令，粘贴到该设备的用户视图（<Huawei>）后回车即可。"
        "命令块以 system-view 开头、return 结尾；执行完记得 save。"
    )
    note_run.font.size = Pt(9)
    note_run.font.color.rgb = RGBColor(0x99, 0x33, 0x00)

    doc.add_heading("拓扑概览", level=1)
    doc.add_heading("设备", level=2)
    _device_table(doc, topo, req)

    doc.add_heading("连线", level=2)
    _link_table(doc, topo)

    _address_section(doc, topo, req)

    if topo.warnings:
        doc.add_heading("解析告警", level=2)
        for warning in topo.warnings:
            doc.add_paragraph(warning, style="List Bullet")

    used: set[str] = set()
    for device in topo.devices:
        commands = configs.get(device.name)
        if not commands:
            continue
        used.add(device.name)
        spec = req.devices.get(device.name)
        dept = spec.department if spec and spec.department else None
        title_name = f"{device.name}（{dept}）" if dept else device.name
        doc.add_page_break()
        doc.add_heading(f"{title_name}　{device.model}", level=1)

        layer_name = LAYER_NAMES.get(spec.layer, spec.layer) if spec and spec.layer else None
        info_text = (
            f"类型：{'路由器' if device.device_type == 'router' else '交换机' if device.device_type == 'switch' else '终端'}"
            f"　接口：{len(device.ports)} 个"
        )
        if layer_name:
            info_text += f"　层级：{layer_name}"
        info = doc.add_paragraph(info_text)
        info.runs[0].font.size = Pt(9)

        doc.add_paragraph("把下面整段命令复制到该设备：")
        _command_block(doc, to_safe_paste(commands))

        tail = doc.add_paragraph("执行完成后保存配置：save")
        tail.runs[0].font.size = Pt(9)

    idle = [d for d in topo.devices if d.name not in used]
    if idle:
        doc.add_heading("未生成配置的设备", level=1)
        doc.add_paragraph(
            "以下设备在 requirements.yaml 中没有配置意图，因此没有命令："
            + "、".join(d.name for d in idle)
        )

    return doc


def write_docx(
    topo: Topology, req: Requirements, configs: dict[str, list[str]], path: str | Path
) -> Path:
    """构建并写盘，返回实际路径。"""

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    build_document(topo, req, configs).save(path)
    return path
