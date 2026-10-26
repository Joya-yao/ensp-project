"""地址表 Excel：在 Excel 里改地址，生成命令时以它为准。

两页签：
- 「设备地址」设备 / 接口 / IP / 掩码 / VLAN / 网关 / 说明
- 「终端地址」终端 / 接在 / VLAN / IP / 掩码 / 网关

没装 openpyxl 时本模块不可用，调用方要检查 :func:`available`。
"""

from __future__ import annotations

import ipaddress
from dataclasses import replace
from pathlib import Path

from labgen.requirements import Requirements
from labgen.topo import Topology

__all__ = [
    "available",
    "write_address_xlsx",
    "read_address_xlsx",
    "apply_overrides",
    "DEVICE_SHEET",
    "TERMINAL_SHEET",
]

DEVICE_SHEET = "设备地址"
TERMINAL_SHEET = "终端地址"
DEVICE_HEADERS = ("设备", "接口", "IP", "掩码", "VLAN", "网关", "说明")
TERMINAL_HEADERS = ("终端", "接在", "VLAN", "IP", "掩码", "网关")


def available() -> bool:
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        return False
    return True


def _rows_from_requirements(topo: Topology, req: Requirements):
    from labgen.addresses import device_rows, terminal_rows

    devices = [list(row) for row in device_rows(req)]
    terminals = [list(row) for row in terminal_rows(topo, req)]
    return devices, terminals


def _rows_from_skeleton(topo: Topology):
    devices = []
    connected = {
        (endpoint.device, endpoint.port.lower())
        for link in topo.links
        for endpoint in (link.src, link.dst)
        if endpoint.port
    }
    for device in topo.devices:
        if device.device_type == "pc":
            continue
        for port in device.ports:
            if (device.name, port.name.lower()) in connected:
                devices.append([device.name, port.name, "", "", "", "", "—"])
    terminals = []
    for device in topo.devices:
        if device.device_type != "pc":
            continue
        where = "未连线"
        for link in topo.links_of(device.name):
            other = link.dst if link.src.device == device.name else link.src
            if other.device != device.name:
                where = f"{other.device}:{other.port}"
                break
        terminals.append([device.name, where, "", "", "", ""])
    return devices, terminals


def write_address_xlsx(
    path: str | Path, topo: Topology, req: Requirements | None = None
) -> Path:
    """写地址表工作簿；有 YAML 时填真实值，没有就写空表待填。"""

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    devices, terminals = (
        _rows_from_requirements(topo, req) if req is not None else _rows_from_skeleton(topo)
    )

    book = Workbook()
    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDE6F5")

    def fill(sheet, headers, rows, widths):
        sheet.append(list(headers))
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for row in rows:
            sheet.append(row)
        for letter, width in zip("ABCDEFG", widths):
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"

    first = book.active
    first.title = DEVICE_SHEET
    fill(first, DEVICE_HEADERS, devices, (12, 26, 16, 16, 8, 16, 24))

    second = book.create_sheet(TERMINAL_SHEET)
    fill(second, TERMINAL_HEADERS, terminals, (10, 26, 8, 16, 16, 16))

    guide = book.create_sheet("说明")
    guide.column_dimensions["A"].width = 88
    for line in (
        "改完这里保存，再在网页上点「生成配置」，生成时会以这张表为准。",
        "「设备地址」：改 IP / 掩码 / VLAN 会覆盖 requirements.yaml 里同一设备同一接口的配置。",
        "「终端地址」：改 VLAN 会覆盖该终端所连交换机端口的 access VLAN。",
        "留空或写 ? 表示不改，沿用 YAML。",
    ):
        guide.append([line])

    book.save(path)
    return path


def _mask_to_prefix(mask: str) -> int:
    return ipaddress.IPv4Network(f"0.0.0.0/{mask}").prefixlen


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def read_address_xlsx(path: str | Path) -> dict[str, list[dict]]:
    """读回 Excel，返回 {"devices": [...], "terminals": [...]}，跳过空行。"""

    from openpyxl import load_workbook

    book = load_workbook(Path(path), data_only=True)
    result: dict[str, list[dict]] = {"devices": [], "terminals": []}

    if DEVICE_SHEET in book.sheetnames:
        sheet = book[DEVICE_SHEET]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = [_clean(cell) for cell in (list(row) + [None] * len(DEVICE_HEADERS))[:7]]
            if not values[0] or not values[1]:
                continue
            result["devices"].append(dict(zip(DEVICE_HEADERS, values)))

    if TERMINAL_SHEET in book.sheetnames:
        sheet = book[TERMINAL_SHEET]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            values = [_clean(cell) for cell in (list(row) + [None] * len(TERMINAL_HEADERS))[:6]]
            if not values[0]:
                continue
            result["terminals"].append(dict(zip(TERMINAL_HEADERS, values)))

    return result


def _is_blank(value: str) -> bool:
    return value in ("", "?", "？", "—", "-")


def apply_overrides(req: Requirements, sheet: dict[str, list[dict]]) -> Requirements:
    """把 Excel 里改过的地址/VLAN 覆盖进 Requirements；空值表示不改。"""

    devices = dict(req.devices)

    for row in sheet.get("devices", []):
        name = row.get("设备", "")
        interface = row.get("接口", "")
        spec = devices.get(name)
        if spec is None:
            continue

        ip_text = row.get("IP", "")
        mask_text = row.get("掩码", "")
        vlan_text = row.get("VLAN", "")

        new_ip = None
        if not _is_blank(ip_text):
            prefix = None
            if not _is_blank(mask_text):
                try:
                    prefix = _mask_to_prefix(mask_text)
                except ValueError:
                    prefix = None
            if prefix is None and "/" in ip_text:
                new_ip = ip_text
            elif prefix is not None:
                new_ip = f"{ip_text}/{prefix}"
            else:
                new_ip = ip_text

        new_vlan = None
        if not _is_blank(vlan_text):
            try:
                new_vlan = int(vlan_text)
            except ValueError:
                new_vlan = None

        if new_ip is None and new_vlan is None:
            continue

        changed = False
        interfaces = []
        for iface in spec.interfaces:
            if iface.name.lower() != interface.lower():
                interfaces.append(iface)
                continue
            interfaces.append(
                replace(
                    iface,
                    ip=new_ip if new_ip is not None else iface.ip,
                    vlan=new_vlan if new_vlan is not None else iface.vlan,
                )
            )
            changed = True
        if changed:
            devices[name] = replace(spec, interfaces=tuple(interfaces))

    for row in sheet.get("terminals", []):
        where = row.get("接在", "")
        vlan_text = row.get("VLAN", "")
        if not where or ":" not in where or _is_blank(vlan_text):
            continue
        device_name, port = where.split(":", 1)
        try:
            vlan = int(vlan_text)
        except ValueError:
            continue
        spec = devices.get(device_name)
        if spec is None:
            continue

        interfaces = []
        changed = False
        for iface in spec.interfaces:
            if iface.name.lower() == port.lower():
                interfaces.append(replace(iface, vlan=vlan))
                changed = True
            else:
                interfaces.append(iface)
        if changed:
            devices[device_name] = replace(spec, interfaces=tuple(interfaces))

    return replace(req, devices=devices)
